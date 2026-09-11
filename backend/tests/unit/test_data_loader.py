import pytest
import openpyxl
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parents[2]
FULL_BOOK = BASE_DIR / 'VSM_Fibre-Q_2024.xlsx'
MAIN_BOOK = BASE_DIR / 'Book1.xlsx'

pytestmark = pytest.mark.unit

ALIASES = {
    'fibre prep fiber': 'Fibre Prep',
    'fibre prep': 'Fibre Prep',
    'fibre prep fiber ': 'Fibre Prep',
    'cell build': 'Cell Build',
    'package assembly': 'Package Assembly',
    'ca': 'Collimator Assembly',
    'collimator assembly': 'Collimator Assembly',
    'ctia': 'CTIA',
}


def test_excel_data_sheet_loading():
    """Verify that openpyxl can read raw stage and process data from the manufacturing workbook."""
    target_book = FULL_BOOK if FULL_BOOK.exists() else MAIN_BOOK
    if not target_book.exists():
        pytest.skip(f"Test Excel workbook not found at {target_book}")

    wb = openpyxl.load_workbook(target_book, data_only=True, read_only=True)
    assert 'Data' in wb.sheetnames, f"'Data' sheet must be in workbook: {wb.sheetnames}"

    ws = wb['Data']
    data = {}
    stage = None
    count = 0

    for row in ws.iter_rows(min_row=1, max_col=10, values_only=True):
        if len(row) < 10:
            continue
        a, b, c, mt, ut, st, op, ct, scrap, rft = row[:10]
        if a is not None:
            key = ' '.join(str(a).strip().lower().split())
            if key in ALIASES or any(k in key for k in ['prep', 'build', 'assembly', 'ca', 'ctia']):
                stage = ALIASES.get(key, str(a).strip())
        if stage and c is not None and ct is not None and str(c).strip().lower() not in ['process', 'name']:
            data.setdefault(stage, []).append({
                'name': str(c).strip(),
                'mt': float(mt or 0),
                'ut': float(ut or 0),
                'st': float(st or 0),
                'op': int(op or 0),
                'ct': float(ct or 0),
                'scrap': float(scrap or 0),
                'rft': float(rft or 0),
                'stage': stage,
            })
            count += 1

    assert count > 0, "Should parse at least one process from the Data sheet"
    assert len(data) >= 3, f"Expected at least 3 stages with processes, got: {list(data.keys())}"

    for stg_name, procs in data.items():
        assert len(procs) > 0, f"Stage {stg_name} must have at least 1 process"
        for p in procs:
            assert p['name'] != "", "Process name must not be empty"
            assert p['ct'] >= 0, f"Cycle time must be non-negative: {p['ct']}"
