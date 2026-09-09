import openpyxl
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent
FULL_BOOK = BASE_DIR / 'VSM_Fibre-Q_2024.xlsx'
MAIN_BOOK = BASE_DIR / 'Book1.xlsx'

ALIASES = {
    'fibre prep fiber': 'Fibre Prep',
    'fibre prep': 'Fibre Prep',
    'fibre prep fiber ': 'Fibre Prep',
    'cell build': 'Cell Build',
    'package assembly': 'Package Assembly',
    'ca': 'Collimator Assembly',
    'collimator assembly': 'Collimator Assembly',
    'ctia': 'CTIA',
}

def test_inspect():
    target_book = FULL_BOOK if FULL_BOOK.exists() else MAIN_BOOK
    print("Reading book:", target_book)
    wb = openpyxl.load_workbook(target_book, data_only=True, read_only=True)
    print("Sheets:", wb.sheetnames)
    
    if 'Data' in wb.sheetnames:
        ws = wb['Data']
        data = {}
        stage = None
        count = 0
        for row in ws.iter_rows(min_row=1, max_col=10, values_only=True):
            if len(row) < 10:
                continue
            a, b, c, mt, ut, st, op, ct, scrap, rft = row[:10]
            if a is not None:
                key = ' '.join(str(a).strip().lower().split())
                if key in ALIASES or any(k in key for k in ['prep', 'build', 'assembly', 'ca', 'ctia']):
                    stage = ALIASES.get(key, str(a).strip())
            if stage and c is not None and ct is not None and str(c).strip().lower() not in ['process', 'name']:
                data.setdefault(stage, []).append({
                    'name': str(c).strip(),
                    'mt': float(mt or 0),
                    'ut': float(ut or 0),
                    'st': float(st or 0),
                    'op': int(op or 0),
                    'ct': float(ct or 0),
                    'scrap': float(scrap or 0),
                    'rft': float(rft or 0),
                    'stage': stage,
                })
                count += 1
        print(f"Total processes parsed: {count}")
        for stg, procs in data.items():
            print(f"  Stage: '{stg}' -> {len(procs)} processes")
            for p in procs[:3]:
                print(f"    - {p['name']} (C/T={p['ct']}, M/T={p['mt']})")

if __name__ == '__main__':
    test_inspect()
