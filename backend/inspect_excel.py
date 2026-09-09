import openpyxl
import pandas as pd
from app.parsers.excel_parser import ExcelVSMParser

files = [
    "VSM_Fibre-Q_2024.xlsx",
    "VSM_Fibre-Q_2024 - Copy.xlsx"
]

parser = ExcelVSMParser()

for fname in files:
    print(f"\n==========================================")
    print(f"INSPECTING FILE: {fname}")
    print(f"==========================================")
    try:
        with open(fname, "rb") as f:
            content = f.read()

        wb = openpyxl.load_workbook(fname, data_only=True)
        print(f"Sheet names in {fname}: {wb.sheetnames}")

        for sname in wb.sheetnames:
            ws = wb[sname]
            print(f"\n--- Sheet: '{sname}' (max_row={ws.max_row}, max_column={ws.max_column}) ---")
            for r_idx, row in enumerate(ws.iter_rows(max_row=15, values_only=True), 1):
                non_empty = [f"Col{i+1}:{v}" for i, v in enumerate(row) if v is not None]
                if non_empty:
                    print(f"Row {r_idx:2d}: {non_empty[:8]}")

        res = parser.parse(content)
        print(f"\nParser Result for {fname}:")
        print(f"  Success: {res.success}")
        if not res.success:
            print(f"  Error Detail: {res.error.model_dump() if hasattr(res.error, 'model_dump') else res.error}")
        else:
            print(f"  Stats: {res.stats}")
            if res.vsm:
                print(f"  Stages Count: {len(res.vsm.stages)}")
                for stg in res.vsm.stages:
                    print(f"    - Stage '{stg.name}': starting_inv={stg.starting_inventory}, inv={stg.inventory}, processes={len(stg.processes)}")

    except Exception as e:
        print(f"Error inspecting {fname}: {e}")
        import traceback
        traceback.print_exc()
